from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

from src.models import CandidateProfile, GeneratedArtifacts, JobPosting, TailoredResumeBundle
from src.utils.config import ResumeConfig
from src.utils.files import ensure_dir, ensure_parent_dir
from src.utils.text import slugify


class ATSSafeResumeDocxGenerator:
    def __init__(self, config: ResumeConfig, config_dir: Path) -> None:
        self.config = config
        self.config_dir = config_dir

    def write_bundle(
        self,
        *,
        profile: CandidateProfile,
        job: JobPosting,
        bundle: TailoredResumeBundle,
    ) -> GeneratedArtifacts:
        docx_dir = (self.config_dir / self.config.docx_output_dir).resolve()
        cover_dir = (self.config_dir / self.config.cover_letter_output_dir).resolve()
        pitch_dir = (self.config_dir / self.config.pitch_output_dir).resolve()
        ensure_dir(docx_dir)
        ensure_dir(cover_dir)
        ensure_dir(pitch_dir)

        file_stem = slugify(f"{job.source}-{job.company}-{job.title}-{job.job_id}")[:100]
        docx_path = docx_dir / f"{file_stem}.docx"
        resume_text_path = docx_dir / f"{file_stem}.txt"
        cover_path = cover_dir / f"{file_stem}.txt"
        pitch_path = pitch_dir / f"{file_stem}.txt"

        document = Document()
        self._set_default_style(document)
        document.add_heading(profile.name, level=0)
        document.add_paragraph(
            " | ".join(
                filter(
                    None,
                    [
                        profile.location,
                        profile.email,
                        profile.phone,
                        str(profile.linkedin or ""),
                        str(profile.portfolio or ""),
                    ],
                )
            )
        )
        document.add_paragraph(bundle.headline)
        self._add_section(document, "Professional Summary", bundle.summary)
        self._add_section(document, "Core Skills", ", ".join(bundle.skills))

        document.add_heading("Experience", level=1)
        for item in bundle.experience:
            document.add_paragraph(f"{item.title} | {item.company} | {item.dates}")
            if item.location:
                document.add_paragraph(item.location)
            for bullet in item.bullets:
                document.add_paragraph(bullet, style="List Bullet")

        if bundle.education:
            document.add_heading("Education", level=1)
            for item in bundle.education:
                document.add_paragraph(f"{item.degree} | {item.institution} | {item.year}")

        if bundle.projects:
            document.add_heading("Projects", level=1)
            for item in bundle.projects:
                document.add_paragraph(f"{item.name}: {item.description}")

        if bundle.certifications:
            self._add_section(document, "Certifications", "\n".join(bundle.certifications))

        document.save(docx_path)
        self._write_text_file(resume_text_path, bundle.resume_text)
        self._write_text_file(cover_path, bundle.cover_letter)
        self._write_text_file(pitch_path, bundle.recruiter_pitch)

        return GeneratedArtifacts(
            resume_docx_path=str(docx_path),
            resume_text_path=str(resume_text_path),
            cover_letter_path=str(cover_path),
            recruiter_pitch_path=str(pitch_path),
        )

    def _set_default_style(self, document: Document) -> None:
        style = document.styles["Normal"]
        style.font.name = self.config.font_name
        style.font.size = Pt(self.config.font_size)

    @staticmethod
    def _add_section(document: Document, title: str, body: str) -> None:
        document.add_heading(title, level=1)
        document.add_paragraph(body)

    @staticmethod
    def _write_text_file(path: Path, contents: str) -> None:
        ensure_parent_dir(path)
        path.write_text(contents, encoding="utf-8")

